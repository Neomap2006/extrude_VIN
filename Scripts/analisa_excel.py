"""
VERIFIKASI MANUAL: cek nilai actual tiap kolom di sheet DL row pertama.
"""
import openpyxl
import sys
sys.stdout.reconfigure(encoding='utf-8')

file = r"folder excel\Masterlist_Digitalisasi_Combined_Fix.xlsx"
wb   = openpyxl.load_workbook(file, data_only=True)
ws   = wb['DL']

# Ambil baris data pertama yang valid (row 6)
ROW = 6
print(f"Verifikasi sheet DL, row {ROW}:\n")
for col in range(1, 54):
    # header dari row 1 dan row 2 (multi-header)
    h1 = ws.cell(row=1, column=col).value
    h2 = ws.cell(row=2, column=col).value
    val= ws.cell(row=ROW, column=col).value
    if val is not None:
        header = f"{h1 or ''} / {h2 or ''}".strip(' /')
        print(f"  col {col:3}: [{header[:40]:40}] = {repr(str(val)[:50])}")

wb.close()
