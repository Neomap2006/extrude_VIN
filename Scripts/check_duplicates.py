import pyodbc
import openpyxl

driver = [d for d in pyodbc.drivers() if 'SQL Server' in d][-1]
conn = pyodbc.connect(f'DRIVER={{{driver}}};SERVER=10.14.149.34,1433;DATABASE=prd_extrude_hose;UID=usrvelasto;PWD=H1s@na2025!!;TrustServerCertificate=yes;Encrypt=no;')
cur = conn.cursor()

# Get total newly imported
cur.execute("""
    SELECT COUNT(*) FROM SpsNoDocs 
    WHERE DocumentNumber LIKE 'SOP/PROD/HOSE/SPS/1/%' 
       OR DocumentNumber LIKE 'SOP/PROD/HOSE/SPS/2/%' 
       OR DocumentNumber LIKE 'SOP/PROD/HOSE/SPS/3/%'
       OR DocumentNumber LIKE 'SOP/PROD/HOSE/SPS/24/%'
""")
total = cur.fetchone()[0]
print(f'Total data yang berhasil masuk ke Database: {total} dokumen')

# Check duplicates in DB (SPS records)
cur.execute("""
    SELECT DocumentNumber, COUNT(*) as cnt 
    FROM SpsNoDocs 
    WHERE DocumentNumber LIKE 'SOP/PROD/HOSE/SPS/%' 
    GROUP BY DocumentNumber 
    HAVING COUNT(*) > 1
""")
dupes_db = cur.fetchall()
if dupes_db:
    print(f"\nDitemukan duplikat DocumentNumber di Database ({len(dupes_db)}):")
    for r in dupes_db:
        print(f"  {r[0]} -> {r[1]} kali")
else:
    print("\nTidak ada duplikat DocumentNumber di Database (semua unik).")

# Now check the Excel file to see if there were duplicates in the source
wb = openpyxl.load_workbook(r'folder excel\Masterlist_Digitalisasi_Combined_Fix.xlsx', data_only=True)
all_docs = {}
dupes_excel = []
for sheet in ['DL', 'CHS 2', 'CHS 3']:
    ws = wb[sheet]
    start_row = 6 if sheet != 'CHS 3' else 5
    for r in range(start_row, ws.max_row + 1):
        doc_val = ws.cell(row=r, column=4).value
        if doc_val and 'SOP/PROD/HOSE/SPS/' in str(doc_val):
            doc = str(doc_val).strip()
            if doc in all_docs:
                dupes_excel.append((doc, all_docs[doc], sheet))
            else:
                all_docs[doc] = sheet

if dupes_excel:
    print(f"\nDitemukan duplikat DocumentNumber di dalam file Excel ({len(dupes_excel)}):")
    for doc, prev_sheet, curr_sheet in dupes_excel:
        print(f"  {doc} ada di sheet '{prev_sheet}' dan juga di sheet '{curr_sheet}'")
else:
    print("\nTidak ada duplikat DocumentNumber di dalam file Excel.")
