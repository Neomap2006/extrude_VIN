import openpyxl
import sqlite3
import glob
import sys
sys.stdout.reconfigure(encoding='utf-8')

# Buka backup SQLite
files = sorted(glob.glob(r'Scripts\backup_velasto_*.sqlite'), reverse=True)
db_file = files[0]
conn = sqlite3.connect(db_file)
cur  = conn.cursor()

# Ambil data yang sudah ada di DB
cur.execute('SELECT DocumentNumber FROM SpsNoDocs')
existing_docs = {r[0].strip() for r in cur.fetchall() if r[0]}

cur.execute('SELECT ItemList, DocumentNumber FROM SpsItemLists')
existing_items = {(r[0].strip(), r[1].strip()) for r in cur.fetchall() if r[0] and r[1]}
conn.close()

# Baca Excel
excel_file = r"folder excel\Masterlist_Digitalisasi_Combined_Fix.xlsx"
wb = openpyxl.load_workbook(excel_file, data_only=True)

sheet_config = {
    'DL':    {'doc_col': 4, 'item_col': 1, 'start_row': 6},
    'CHS 2': {'doc_col': 4, 'item_col': 1, 'start_row': 6},
    'CHS 3': {'doc_col': 4, 'item_col': 1, 'start_row': 5},
}

new_docs = set()
new_items = set()

for sheet_name, cfg in sheet_config.items():
    ws = wb[sheet_name]
    for row_idx in range(cfg['start_row'], ws.max_row + 1):
        item_val = ws.cell(row=row_idx, column=cfg['item_col']).value
        doc_val  = ws.cell(row=row_idx, column=cfg['doc_col']).value
        
        if not doc_val or not str(doc_val).startswith('SOP/'): continue
        if not item_val or str(item_val) in ['1','2','3','Item','No','ID']: continue
        
        item = str(item_val).strip()
        doc = str(doc_val).strip()
        
        if doc not in existing_docs:
            new_docs.add(doc)
            
        if (item, doc) not in existing_items:
            new_items.add((item, doc))

wb.close()

print("=== DAFTAR NO DOCUMENT (SpsNoDoc) BARU ===")
print(f"Total: {len(new_docs)} dokumen")
for d in sorted(new_docs):
    print(f" - {d}")

print("\n=== DAFTAR ITEM LIST (SpsItemList) BARU ===")
print(f"Total: {len(new_items)} item (beserta No Document pasangannya)")
for item, doc in sorted(new_items, key=lambda x: (x[1], x[0])):
    print(f" - Item: {item:<10} | Doc: {doc}")
