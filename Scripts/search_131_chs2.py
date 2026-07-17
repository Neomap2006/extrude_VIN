import openpyxl

file_path = r'folder excel\Masterlist SPS CHS 2 Layer DIG.xlsx'
print(f"Mencari 131 di {file_path}")
wb = openpyxl.load_workbook(file_path, data_only=True)
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            val = ws.cell(row=r, column=c).value
            if val and '131' in str(val):
                print(f"Sheet: {sheet_name}, Row: {r}, Col: {c}, Value: {repr(val)}")
