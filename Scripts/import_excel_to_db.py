"""
====================================================================
IMPORT EXCEL → SpsNoDoc & SpsItemList
====================================================================
SAFE: Hanya INSERT data baru. Tidak ada UPDATE, DELETE, atau DROP.
Jika DocumentNumber sudah ada di DB → SKIP (tidak diimport ulang).
====================================================================
"""
import openpyxl
import pyodbc
import sys
sys.stdout.reconfigure(encoding='utf-8')

# ── KONFIGURASI KONEKSI ──────────────────────────────────────────
SQL_SERVER   = "10.14.149.34,1433"
SQL_DB       = "prd_extrude_hose"
SQL_USER     = "usrvelasto"
SQL_PASSWORD = "H1s@na2025!!"

EXCEL_FILE   = r"folder excel\Masterlist_Digitalisasi_Combined_Fix.xlsx"

DRY_RUN = False  # ← Set False untuk benar-benar insert ke DB

# ── MAPPING KOLOM EXCEL → SpsNoDoc ──────────────────────────────
# Format: 'field_db': kolom_excel (1-indexed)
# Verified dari scan kolom per sheet

# Sheet DL (Double Layer) — VERIFIED dari scan aktual row 6
MAPPING_DL = {
    # IDENTITAS
    'No':             3,   # col 3:  No urut
    'DocumentNumber': 4,   # col 4:  No Document (PK)
    'RevisionNumber': 7,   # col 7:  No Rev
    'Customer':       8,   # col 8:  Customer
    'RevisionDate':   9,   # col 9:  Revision Date
    'Formulasi':      10,  # col 10: Formulasi
    'HoseType':       11,  # col 11: Hose Type
    'Dimensi':        12,  # col 12: Dimensi
    'Material':       13,  # col 13: Material TSM
    # MATERIAL
    'InnerTube':      14,  # col 14: Inner / #tube
    'OuterCover':     15,  # col 15: Outer / #cover
    # MiddleTube tidak ada di DL
    'UseLimitsInner': 16,  # col 16: Use limits #Inner
    'UseLimitsOuter': 17,  # col 17: Use limits #Outer
    'Nipple':         18,  # col 18: Nipple
    'TubeDie':        19,  # col 19: Tube Die
    'CoverDie':       20,  # col 20: Cover Die
    # Mesh Screen — DL hanya 2 kolom (tidak ada dim terpisah)
    'MeshScreen1':    21,  # col 21: Mesh screen #material 1
    'MeshScreen2':    22,  # col 22: Mesh screen #material 2
    # Temperature #material 1 (inner)
    'HeadTemp1':      23,  # col 23: Head Temp. #material 1
    'HeadTemp2':      24,  # col 24: Head Temp. #material 2
    'Cylinder1_1':    25,  # col 25: Cylinder 1 #material 1
    'Cylinder1_2':    26,  # col 26: Cylinder 1 #material 2
    'Cylinder2_1':    27,  # col 27: Cylinder 2 #material 1
    'Cylinder2_2':    28,  # col 28: Cylinder 2 #material 2
    'Feed1':          29,  # col 29: FEED #material 1
    'Feed2':          30,  # col 30: FEED #material 2
    'ScrewTemp1':     31,  # col 31: Screw Temp. #material 1
    'ScrewTemp2':     32,  # col 32: Screw Temp. #material 2
    'ScrewSpeed1':    33,  # col 33: Screw Speed #material 1
    'ScrewSpeed2':    34,  # col 34: Screw Speed #material 2
    'Pressure1':      35,  # col 35: Pressure #material 1
    'Pressure2':      36,  # col 36: Pressure #material 2
    # col 37: Am meter — kosong di DL (tidak ada nilai)
    'OdSensor':       38,  # col 38: OD sensor = 7.3±0.1
    'MarkingSort':    39,  # col 39: Marking Sort
    'TextMarkingMaterial': 40,  # col 40: Text marking mt'l
    'MarkingColour':  41,  # col 41: Marking Colour
    'ChillerWaterTemp': 42, # col 42: Chiller water temp.
    'CuttingSpeed':   43,  # col 43: Cutting speed
    'TakeUpConveyorSpeed': 44,  # col 44: Take up conveyor speed
    'ToleranceInner': 45,  # col 45: ± Inner
    'ToleranceOuter': 46,  # col 46: ± Outer
    'TebalInner':     47,  # col 47: Tebal Inner
    'TebalOuter':     48,  # col 48: Tebal Outer
    'TebalTotal':     49,  # col 49: Tebal Total
    'SelisihTebal':   50,  # col 50: Selisih tebal
    'MachineCode':    51,  # col 51: MC
}

# Sheet CHS 2 — verified dari scan
MAPPING_CHS2 = {
    'No':             3,
    'DocumentNumber': 4,
    'RevisionNumber': 7,
    'Customer':       8,
    'RevisionDate':   9,
    'Formulasi':      10,
    'HoseType':       11,
    'Dimensi':        12,
    'Material':       13,
    'InnerTube':      14,  # Inner
    'OuterCover':     15,  # Outer
    'UseLimitsInner': 16,
    'UseLimitsOuter': 17,
    'Yarn':           18,
    'PitchYarn':      19,
    'TensionYarnInner': 20,
    'TensionYarnOuter': 21,
    'Nipple':         22,
    'TubeDie':        23,
    'MiddleDie':      24,  # Middle
    'CoverDie':       25,
    'SpacerDie':      26,
    'ADistance':      27,
    'MeshDim1':       28,
    'MeshScreen1':    29,  # Note: row1=Mesh screen, row2=#material 1
    'MeshDim2':       30,
    'MeshScreen2':    31,
    # Temperature #material 1 (inner)
    'HeadTemp1':      32,
    'Cylinder1_1':    33,
    'Cylinder2_1':    34,
    'Cylinder3_1':    35,
    'ScrewTemp1':     36,
    # Temperature #material 2 (outer)
    'HeadTemp2':      37,
    'Cylinder1_2':    38,
    'Cylinder2_2':    39,
    'Cylinder3_2':    40,
    'ScrewTemp2':     41,
    'ScrewSpeed1':    42,
    'ScrewSpeed2':    43,
    'FeedRollRatio1': 44,
    'FeedRollRatio2': 45,
    'Pressure1':      46,
    'Pressure2':      47,
    'CurrentValue':   48,
    'AmMeter':        49,
    'AmMeter2':       50,
    'PresetValue':    51,
    'ControlValue':   52,
    'SpiralPitchSetting': 53,
    'SpiralPitchDisplay': 54,
    'SpiralSpeed':    55,
    'HoseSpeed':      56,
    'UnsmoothSurface':57,
    'MarkingSort':    58,
    'TextMarkingMaterial': 59,
    'MarkingColour':  60,
    'ChillerWaterTemp':61,
    'DancerPosition': 62,
    'CaterpillarGap': 63,
    'TakeUpConveyorSpeed': 64,
    'CoolConveyorSpeed':   65,
    'ConveyorRatio':       66,
    'ToleranceInner': 67,
    'ToleranceOuter': 68,
    'TebalInner':     69,
    'TebalTotal':     70,
    'SelisihTebal':   71,
    # Quality matrix
    'InnerTarget':    72,
    'InnerTol':       73,
    'InnerLCL':       74,
    'InnerMin':       75,
    'InnerUCL':       76,
    'InnerMax':       77,
    'ThickTarget':    78,
    'ThickTol':       79,
    'ThickLCL':       80,
    'ThickMin':       81,
    'ThickUCL':       82,
    'ThickMax':       83,
    'TotalTarget':    84,
    'TotalTol':       85,
    'TotalLCL':       86,
    'TotalMin':       87,
    'TotalUCL':       88,
    'TotalMax':       89,
}

# Sheet CHS 3 — verified dari scan
MAPPING_CHS3 = {
    'No':             3,
    'DocumentNumber': 4,
    'RevisionNumber': 7,
    'Customer':       8,
    'RevisionDate':   9,
    'Formulasi':      10,
    'HoseType':       11,
    'Dimensi':        12,
    'Material':       13,
    'InnerTube':      14,  # #tube1
    'MiddleTube':     15,  # #tube2
    'OuterCover':     16,  # #cover
    'UseLimitsInner': 17,
    'UseLimitsMiddle':18,
    'UseLimitsOuter': 19,
    'Yarn':           20,
    'TensionYarnInner': 21,
    'TensionYarnOuter': 22,
    'Nipple':         23,
    'TubeDie':        24,
    'MiddleDie':      25,
    'CoverDie':       26,
    'SpacerDie':      27,
    'ADistance':      28,
    'MeshDim1':       29,
    'MeshScreen1':    30,
    'MeshDim2':       31,
    'MeshScreen2':    32,
    'MeshDim3':       33,
    'MeshScreen3':    34,
    # Temperature #material 1 (inner)
    'HeadTemp1':      35,
    'Cylinder1_1':    36,
    'Cylinder2_1':    37,
    'Cylinder3_1':    38,
    'ScrewTemp1':     39,
    # Temperature #material 2 (middle)
    'HeadTemp2':      40,
    'Cylinder1_2':    41,
    'Cylinder2_2':    42,
    'Cylinder3_2':    43,
    'ScrewTemp2':     44,
    # Temperature #material 3 (outer)
    'HeadTemp3':      45,
    'Cylinder1_3':    46,
    'Cylinder2_3':    47,
    'Cylinder3_3':    48,
    'ScrewTemp3':     49,
    'ScrewSpeed1':    50,
    'ScrewSpeed2':    51,
    'ScrewSpeed3':    52,
    'FeedRollRatio1': 53,
    'FeedRollRatio2': 54,
    'FeedRollRatio3': 55,
    'Pressure1':      56,
    'Pressure2':      57,
    'Pressure3':      58,
    'AmMeter':        59,
    'AmMeter2':       60,
    'AmMeter3':       61,
    'PresetValue':    62,
    'ControlValue':   63,
    'SpiralPitchSetting': 64,
    'SpiralPitchDisplay': 65,
    'SpiralSpeed':    66,
    'HoseSpeed':      67,
    'UnsmoothSurface':68,
    'MarkingSort':    69,
    'TextMarkingMaterial': 70,
    'MarkingColour':  71,
    'ChillerWaterTemp': 72,
    'CaterpillarGap': 73,
    'TakeUpConveyorSpeed': 74,
    'CoolConveyorSpeed':   75,
    'CoolConveyorSpeed2':  76,
    'ConveyorRatio':       77,
    'ToleranceInner': 78,
    'ToleranceOuter': 79,
    'TebalInner':     80,
    'TebalInnerMiddle': 81,
    'TebalTotal':     82,
    'SelisihTebal':   83,
    # Quality matrix (col 84-107)
    'InnerTarget':    84,
    'InnerTol':       85,
    'InnerLCL':       86,
    'InnerMin':       87,
    'InnerUCL':       88,
    'InnerMax':       89,
    'ThickTarget':    90,
    'ThickTol':       91,
    'ThickLCL':       92,
    'ThickMin':       93,
    'ThickUCL':       94,
    'ThickMax':       95,
    'InnerMidTarget': 96,
    'InnerMidTol':    97,
    'InnerMidLCL':    98,
    'InnerMidMin':    99,
    'InnerMidUCL':    100,
    'InnerMidMax':    101,
    'TotalTarget':    102,
    'TotalTol':       103,
    'TotalLCL':       104,
    'TotalMin':       105,
    'TotalUCL':       106,
    'TotalMax':       107,
}

SHEET_CONFIG = {
    'DL':    {'mapping': MAPPING_DL,   'data_start_row': 6,  'doc_col': 4, 'item_col': 1},
    'CHS 2': {'mapping': MAPPING_CHS2, 'data_start_row': 6,  'doc_col': 4, 'item_col': 1},
    'CHS 3': {'mapping': MAPPING_CHS3, 'data_start_row': 5,  'doc_col': 4, 'item_col': 1},
}

def v(ws, row, col):
    """Ambil nilai sel, return string bersih atau None."""
    if col is None:
        return None
    val = ws.cell(row=row, column=col).value
    if val is None:
        return None
    s = str(val).strip()
    return s if s and s != '-' else None

def get_odbc_driver():
    available = [d for d in pyodbc.drivers() if 'SQL Server' in d]
    for p in ['ODBC Driver 18 for SQL Server', 'ODBC Driver 17 for SQL Server', 'SQL Server']:
        if p in available:
            return p
    return available[-1] if available else None

def run_import():
    print('='*70)
    print('  IMPORT EXCEL → SpsNoDoc & SpsItemList')
    mode = 'DRY RUN (tidak ada perubahan ke DB)' if DRY_RUN else '*** LIVE INSERT ke SQL Server ***'
    print(f'  Mode   : {mode}')
    print('='*70)

    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)

    # Kumpulkan semua data terlebih dahulu
    all_docs  = {}   # DocumentNumber → {field: value}
    all_items = []   # list of (ItemList, DocumentNumber)

    for sheet_name, cfg in SHEET_CONFIG.items():
        ws      = wb[sheet_name]
        mapping = cfg['mapping']
        doc_col = cfg['doc_col']
        item_col= cfg['item_col']
        start   = cfg['data_start_row']

        print(f'\nMembaca sheet: {sheet_name}')
        sheet_docs_found = 0
        sheet_items_found = 0

        seen_docs = set()
        for row_idx in range(start, ws.max_row + 1):
            item_val = v(ws, row_idx, item_col)
            doc_val  = v(ws, row_idx, doc_col)

            if not doc_val or not doc_val.startswith('SOP/'):
                continue
            if not item_val or item_val in ['1','2','3','Item','No','ID']:
                continue

            # Catat SpsNoDoc (1 per DocumentNumber)
            if doc_val not in seen_docs:
                seen_docs.add(doc_val)
                row_data = {'IsActive': '1'}
                for field, col_idx in mapping.items():
                    row_data[field] = v(ws, row_idx, col_idx)
                all_docs[doc_val] = row_data
                sheet_docs_found += 1

            # Catat SpsItemList
            all_items.append((item_val, doc_val))
            sheet_items_found += 1

        print(f'  Ditemukan: {sheet_docs_found} dokumen, {sheet_items_found} item')

    wb.close()

    print(f'\nTotal dikumpulkan:')
    print(f'  SpsNoDoc   : {len(all_docs)} dokumen')
    print(f'  SpsItemList: {len(all_items)} item')

    if DRY_RUN:
        print('\n[DRY RUN] Preview 3 dokumen pertama:')
        for i, (doc, data) in enumerate(list(all_docs.items())[:3]):
            print(f'\n  DocumentNumber: {doc}')
            for k, val in data.items():
                if val:
                    print(f'    {k}: {val}')
        print('\n[DRY RUN] 5 item pertama:')
        for item, doc in all_items[:5]:
            print(f'  ItemList={item} → DocumentNumber={doc}')
        print('\n[DRY RUN] Selesai. Set DRY_RUN=False untuk insert ke DB.')
        return

    # ── KONEKSI DB ────────────────────────────────────────────────
    driver = get_odbc_driver()
    conn_str = (f"DRIVER={{{driver}}};SERVER={SQL_SERVER};DATABASE={SQL_DB};"
                f"UID={SQL_USER};PWD={SQL_PASSWORD};TrustServerCertificate=yes;Encrypt=no;")
    conn = pyodbc.connect(conn_str, timeout=15)
    cursor = conn.cursor()

    # Ambil daftar yang sudah ada
    cursor.execute('SELECT DocumentNumber FROM SpsNoDocs')
    existing_docs = {r[0].strip() for r in cursor.fetchall() if r[0]}
    cursor.execute('SELECT ItemList, DocumentNumber FROM SpsItemLists')
    existing_items = {(r[0].strip(), r[1].strip()) for r in cursor.fetchall() if r[0] and r[1]}

    # ── INSERT SpsNoDocs ─────────────────────────────────────────
    print('\nInserting SpsNoDocs...')
    inserted_docs = 0
    skipped_docs  = 0
    errors_docs   = 0

    for doc_num, data in all_docs.items():
        if doc_num in existing_docs:
            skipped_docs += 1
            continue
        fields = list(data.keys())
        placeholders = ', '.join(['?' for _ in fields])
        col_names    = ', '.join([f'[{f}]' for f in fields])
        values       = [data[f] for f in fields]
        sql = f'INSERT INTO SpsNoDocs ({col_names}) VALUES ({placeholders})'
        try:
            cursor.execute(sql, values)
            inserted_docs += 1
            print(f'  + INSERT SpsNoDoc: {doc_num}')
        except Exception as e:
            errors_docs += 1
            print(f'  ! ERROR {doc_num}: {e}')

    conn.commit()
    print(f'  SpsNoDoc: {inserted_docs} inserted, {skipped_docs} skipped, {errors_docs} errors')

    # ── INSERT SpsItemLists ──────────────────────────────────────
    print('\nInserting SpsItemLists...')
    inserted_items = 0
    skipped_items  = 0
    errors_items   = 0

    for item_val, doc_val in all_items:
        key = (item_val, doc_val)
        if key in existing_items:
            skipped_items += 1
            continue
        try:
            cursor.execute(
                'INSERT INTO SpsItemLists (ItemList, DocumentNumber) VALUES (?, ?)',
                (item_val, doc_val)
            )
            inserted_items += 1
        except Exception as e:
            errors_items += 1
            print(f'  ! ERROR item {item_val} / {doc_val}: {e}')

    conn.commit()
    conn.close()
    print(f'  SpsItemList: {inserted_items} inserted, {skipped_items} skipped, {errors_items} errors')

    print('\n' + '='*70)
    print('  IMPORT SELESAI')
    print(f'  SpsNoDoc   : {inserted_docs} baru, {skipped_docs} skip, {errors_docs} error')
    print(f'  SpsItemList: {inserted_items} baru, {skipped_items} skip, {errors_items} error')
    print('='*70)

if __name__ == '__main__':
    run_import()
