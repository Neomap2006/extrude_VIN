import openpyxl

docs = [
    'SOP/PROD/HOSE/EPR/24/10/13', 'SOP/PROD/HOSE/EPR/24/10/16',
    'SOP/PROD/HOSE/EPR/24/10/17', 'SOP/PROD/HOSE/EPR/24/10/22',
    'SOP/PROD/HOSE/EPR/24/10/24', 'SOP/PROD/HOSE/EPR/24/10/27',
    'SOP/PROD/HOSE/EPR/24/10/28', 'SOP/PROD/HOSE/EPR/24/10/37'
]

f = r'folder excel\Masterlist SPS Double Layer.xlsx'
print(f"Membaca {f}...")
wb = openpyxl.load_workbook(f, data_only=True)
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            val = str(ws.cell(row=r, column=c).value or '').strip()
            if val in docs:
                print(f"Found {val} in sheet '{sheet_name}', row {r}, col {c}")
