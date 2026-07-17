import openpyxl

wb = openpyxl.load_workbook(r'folder excel\Masterlist_Digitalisasi_Combined_Fix.xlsx', data_only=True)
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            val = ws.cell(row=r, column=c).value
            if val and '131' in str(val):
                print(f"Sheet: {sheet_name}, Row: {r}, Col: {c}, Value: {repr(val)}")
