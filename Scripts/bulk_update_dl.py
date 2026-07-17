import openpyxl
import pyodbc
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')
SQL_SERVER   = "10.14.149.34,1433"
SQL_DB       = "prd_extrude_hose"
SQL_USER     = "usrvelasto"
SQL_PASSWORD = "H1s@na2025!!"

MAPPING_DL_ORIG = {
    'No': 3, 'DocumentNumber': 4, 'RevisionNumber': 7, 'Customer': 8, 'RevisionDate': 9,
    'Formulasi': 10, 'HoseType': 11, 'Dimensi': 12, 'Material': 13, 'InnerTube': 14,
    'OuterCover': 15, 'UseLimitsInner': 16, 'UseLimitsOuter': 17, 'Nipple': 18, 'TubeDie': 19,
    'CoverDie': 20, 'MeshScreen1': 21, 'MeshScreen2': 22, 'HeadTemp1': 23, 'HeadTemp2': 24,
    'Cylinder1_1': 25, 'Cylinder1_2': 26, 'Cylinder2_1': 27, 'Cylinder2_2': 28, 'Feed1': 29,
    'Feed2': 30, 'ScrewTemp1': 31, 'ScrewTemp2': 32, 'ScrewSpeed1': 33, 'ScrewSpeed2': 34,
    'Pressure1': 35, 'Pressure2': 36, 'OdSensor': 38, 'MarkingSort': 39, 'TextMarkingMaterial': 40,
    'MarkingColour': 41, 'ChillerWaterTemp': 42, 'CuttingSpeed': 43, 'TakeUpConveyorSpeed': 44,
    'ToleranceInner': 45, 'ToleranceOuter': 46, 'TebalInner': 47, 'TebalOuter': 48,
    'TebalTotal': 49, 'SelisihTebal': 50
}
MAPPING_DL = {}
for k, v in MAPPING_DL_ORIG.items():
    if k == 'No': MAPPING_DL[k] = 2
    elif v > 4: MAPPING_DL[k] = v - 2
    else: MAPPING_DL[k] = v

OUTDATED_DOCS = {
    'SOP/PROD/HOSE/EPR/24/10/13': {'file': r'folder excel\Masterlist SPS Double Layer.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_DL, 'machine': 'Double Layer'},
    'SOP/PROD/HOSE/EPR/24/10/16': {'file': r'folder excel\Masterlist SPS Double Layer.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_DL, 'machine': 'Double Layer'},
    'SOP/PROD/HOSE/EPR/24/10/17': {'file': r'folder excel\Masterlist SPS Double Layer.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_DL, 'machine': 'Double Layer'},
    'SOP/PROD/HOSE/EPR/24/10/22': {'file': r'folder excel\Masterlist SPS Double Layer.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_DL, 'machine': 'Double Layer'},
    'SOP/PROD/HOSE/EPR/24/10/24': {'file': r'folder excel\Masterlist SPS Double Layer.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_DL, 'machine': 'Double Layer'},
    'SOP/PROD/HOSE/EPR/24/10/27': {'file': r'folder excel\Masterlist SPS Double Layer.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_DL, 'machine': 'Double Layer'},
    'SOP/PROD/HOSE/EPR/24/10/28': {'file': r'folder excel\Masterlist SPS Double Layer.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_DL, 'machine': 'Double Layer'},
    'SOP/PROD/HOSE/EPR/24/10/37': {'file': r'folder excel\Masterlist SPS Double Layer.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_DL, 'machine': 'Double Layer'},
}

def v(ws, row, col):
    val = ws.cell(row=row, column=col).value
    s = str(val).strip() if val is not None else None
    return s if s and s != '-' else None

print("Mengumpulkan data terbaru dari Excel (Sisa 8 dokumen DL)...")
updates_to_make = {}

f = r'folder excel\Masterlist SPS Double Layer.xlsx'
print(f"Membaca {f}...")
wb = openpyxl.load_workbook(f, data_only=True)

for doc_num, info in OUTDATED_DOCS.items():
    ws = wb[info['sheet']]
    mapping = info['map']
    
    highest_rev = -1
    latest_row_data = None
    
    for r in range(5, ws.max_row + 1):
        # Cek kolom 4 dan kolom 49
        doc_val = v(ws, r, 4)
        if doc_val != doc_num:
            doc_val = v(ws, r, 49)
            
        if doc_val == doc_num:
            rev_col_idx = mapping.get('RevisionNumber', 5)
            rev_val = v(ws, r, rev_col_idx)
            try:
                rev_num = int(rev_val) if rev_val else 0
            except:
                rev_num = 0
            
            if rev_num > highest_rev:
                highest_rev = rev_num
                row_data = {'IsActive': '1', 'Machine': info['machine'], 'MachineCode': info['machine']}
                for field, col_idx in mapping.items():
                    # For DocumentNumber, use the value we found so it doesn't try to fetch from col 4 if it's empty
                    if field == 'DocumentNumber':
                        row_data[field] = doc_num
                    else:
                        row_data[field] = v(ws, r, col_idx)
                latest_row_data = row_data
                
    if latest_row_data:
        updates_to_make[doc_num] = latest_row_data

wb.close()

print(f"Berhasil mengumpulkan data terbaru untuk {len(updates_to_make)} dokumen.")

driver = [d for d in pyodbc.drivers() if 'SQL Server' in d][-1]
conn_str = f"DRIVER={{{driver}}};SERVER={SQL_SERVER};DATABASE={SQL_DB};UID={SQL_USER};PWD={SQL_PASSWORD};TrustServerCertificate=yes;Encrypt=no;"
conn = pyodbc.connect(conn_str)
cur = conn.cursor()

def parse_pm_value(s):
    if not s: return None, None, None
    s = s.replace(',', '.').strip()
    m = re.search(r'([0-9.]+)\s*[±+]\s*([0-9.]+)', s)
    if m:
        base = float(m.group(1))
        tol = float(m.group(2))
        return round(base - tol, 4), round(base, 4), round(base + tol, 4)
    return None, None, None

columns_to_parse = [
    'ToleranceInner', 'ToleranceOuter', 'TebalInner', 'TebalOuter', 'TebalTotal', 
    'TebalInnerMiddle', 'Nipple', 'TubeDie', 'CoverDie', 'MiddleDie', 'SpacerDie', 
    'ADistance', 'MeshDim1', 'MeshDim2', 'MeshDim3', 'HeadTemp1', 'HeadTemp2', 'HeadTemp3',
    'Cylinder1_1', 'Cylinder1_2', 'Cylinder1_3', 'Cylinder2_1', 'Cylinder2_2', 'Cylinder2_3',
    'Cylinder3_1', 'Cylinder3_2', 'Cylinder3_3', 'ScrewTemp1', 'ScrewTemp2', 'ScrewTemp3',
    'ScrewSpeed1', 'ScrewSpeed2', 'ScrewSpeed3', 'Pressure1', 'Pressure2', 'Pressure3',
    'HoseSpeed', 'TakeUpConveyorSpeed', 'ChillerWaterTemp', 'Feed1', 'Feed2', 'Feed3',
    'FeedRollRatio1', 'FeedRollRatio2', 'FeedRollRatio3', 'CaterpillarGap', 'SpiralSpeed',
    'AmMeter', 'AmMeter2', 'AmMeter3', 'PresetValue', 'ControlValue', 'SpiralPitchSetting',
    'SpiralPitchDisplay', 'CoolConveyorSpeed', 'CoolConveyorSpeed2', 'ConveyorRatio',
    'SelisihTebal', 'PitchYarn', 'DancerPosition', 'OdSensor', 'CuttingSpeed', 'ToleranceSpiralPitch'
]

print("Mulai memperbarui database...")
success_count = 0
for doc_num, data in updates_to_make.items():
    try:
        cur.execute("DELETE FROM SpsNoDocs WHERE DocumentNumber = ?", (doc_num,))
        fields = list(data.keys())
        placeholders = ', '.join(['?' for _ in fields])
        col_names = ', '.join([f'[{f}]' for f in fields])
        values = [data[f] for f in fields]
        
        sql = f'INSERT INTO SpsNoDocs ({col_names}) VALUES ({placeholders})'
        cur.execute(sql, values)
        
        updates = []
        for col in columns_to_parse:
            if col in data and data[col]:
                mn, asli, mx = parse_pm_value(data[col])
                if mn is not None:
                    updates.append(f"[{col}_Min] = {mn}")
                    updates.append(f"[{col}_Asli] = {asli}")
                    updates.append(f"[{col}_Max] = {mx}")
                    
        if updates:
            update_sql = f"UPDATE SpsNoDocs SET {', '.join(updates)} WHERE DocumentNumber = ?"
            cur.execute(update_sql, (doc_num,))
            
        success_count += 1
        print(f"  [+] Berhasil update: {doc_num} (Rev {data.get('RevisionNumber')})")
    except Exception as e:
        print(f"  [!] Gagal update {doc_num}: {e}")

conn.commit()
conn.close()
print(f"\\nSelesai! {success_count} dokumen berhasil di-update.")
