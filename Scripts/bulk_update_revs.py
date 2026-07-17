import openpyxl
import pyodbc
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')
SQL_SERVER   = "10.14.149.34,1433"
SQL_DB       = "prd_extrude_hose"
SQL_USER     = "usrvelasto"
SQL_PASSWORD = "H1s@na2025!!"

MAPPING_DL_ORIG = {
    'No': 3, 'DocumentNumber': 4, 'RevisionNumber': 7, 'Customer': 8, 'RevisionDate': 9,
    'Formulasi': 10, 'HoseType': 11, 'Dimensi': 12, 'Material': 13, 'InnerTube': 14,
    'OuterCover': 15, 'UseLimitsInner': 16, 'UseLimitsOuter': 17, 'Nipple': 18, 'TubeDie': 19,
    'CoverDie': 20, 'MeshScreen1': 21, 'MeshScreen2': 22, 'HeadTemp1': 23, 'HeadTemp2': 24,
    'Cylinder1_1': 25, 'Cylinder1_2': 26, 'Cylinder2_1': 27, 'Cylinder2_2': 28, 'Feed1': 29,
    'Feed2': 30, 'ScrewTemp1': 31, 'ScrewTemp2': 32, 'ScrewSpeed1': 33, 'ScrewSpeed2': 34,
    'Pressure1': 35, 'Pressure2': 36, 'OdSensor': 38, 'MarkingSort': 39, 'TextMarkingMaterial': 40,
    'MarkingColour': 41, 'ChillerWaterTemp': 42, 'CuttingSpeed': 43, 'TakeUpConveyorSpeed': 44,
    'ToleranceInner': 45, 'ToleranceOuter': 46, 'TebalInner': 47, 'TebalOuter': 48,
    'TebalTotal': 49, 'SelisihTebal': 50
}
# Adjust for the actual old file where index > 4 is shifted by -2
MAPPING_DL = {}
for k, v in MAPPING_DL_ORIG.items():
    if k == 'No': MAPPING_DL[k] = 2
    elif v > 4: MAPPING_DL[k] = v - 2
    else: MAPPING_DL[k] = v

MAPPING_CHS2 = {
    'No': 3, 'DocumentNumber': 4, 'RevisionNumber': 7, 'Customer': 8, 'RevisionDate': 9,
    'Formulasi': 10, 'HoseType': 11, 'Dimensi': 12, 'Material': 13, 'InnerTube': 14,
    'OuterCover': 15, 'UseLimitsInner': 16, 'UseLimitsOuter': 17, 'Yarn': 18, 'PitchYarn': 19,
    'TensionYarnInner': 20, 'TensionYarnOuter': 21, 'Nipple': 22, 'TubeDie': 23, 'MiddleDie': 24,
    'CoverDie': 25, 'SpacerDie': 26, 'ADistance': 27, 'MeshDim1': 28, 'MeshScreen1': 29,
    'MeshDim2': 30, 'MeshScreen2': 31, 'HeadTemp1': 32, 'Cylinder1_1': 33, 'Cylinder2_1': 34,
    'Cylinder3_1': 35, 'ScrewTemp1': 36, 'HeadTemp2': 37, 'Cylinder1_2': 38, 'Cylinder2_2': 39,
    'Cylinder3_2': 40, 'ScrewTemp2': 41, 'ScrewSpeed1': 42, 'ScrewSpeed2': 43, 'FeedRollRatio1': 44,
    'FeedRollRatio2': 45, 'Pressure1': 46, 'Pressure2': 47, 'CurrentValue': 48, 'AmMeter': 49,
    'AmMeter2': 50, 'PresetValue': 51, 'ControlValue': 52, 'SpiralPitchSetting': 53,
    'SpiralPitchDisplay': 54, 'SpiralSpeed': 55, 'HoseSpeed': 56, 'UnsmoothSurface': 57,
    'MarkingSort': 58, 'TextMarkingMaterial': 59, 'MarkingColour': 60, 'ChillerWaterTemp': 61,
    'DancerPosition': 62, 'CaterpillarGap': 63, 'TakeUpConveyorSpeed': 64, 'CoolConveyorSpeed': 65,
    'ConveyorRatio': 66, 'ToleranceInner': 67, 'ToleranceOuter': 68, 'TebalInner': 69,
    'TebalTotal': 70, 'SelisihTebal': 71, 'InnerTarget': 72, 'InnerTol': 73, 'InnerLCL': 74,
    'InnerMin': 75, 'InnerUCL': 76, 'InnerMax': 77, 'ThickTarget': 78, 'ThickTol': 79,
    'ThickLCL': 80, 'ThickMin': 81, 'ThickUCL': 82, 'ThickMax': 83, 'TotalTarget': 84,
    'TotalTol': 85, 'TotalLCL': 86, 'TotalMin': 87, 'TotalUCL': 88, 'TotalMax': 89,
}

MAPPING_CHS3 = {
    'No': 3, 'DocumentNumber': 4, 'RevisionNumber': 7, 'Customer': 8, 'RevisionDate': 9,
    'Formulasi': 10, 'HoseType': 11, 'Dimensi': 12, 'Material': 13, 'InnerTube': 14,
    'MiddleTube': 15, 'OuterCover': 16, 'UseLimitsInner': 17, 'UseLimitsMiddle': 18,
    'UseLimitsOuter': 19, 'Yarn': 20, 'TensionYarnInner': 21, 'TensionYarnOuter': 22,
    'Nipple': 23, 'TubeDie': 24, 'MiddleDie': 25, 'CoverDie': 26, 'SpacerDie': 27,
    'ADistance': 28, 'MeshDim1': 29, 'MeshScreen1': 30, 'MeshDim2': 31, 'MeshScreen2': 32,
    'MeshDim3': 33, 'MeshScreen3': 34, 'HeadTemp1': 35, 'Cylinder1_1': 36, 'Cylinder2_1': 37,
    'Cylinder3_1': 38, 'ScrewTemp1': 39, 'HeadTemp2': 40, 'Cylinder1_2': 41, 'Cylinder2_2': 42,
    'Cylinder3_2': 43, 'ScrewTemp2': 44, 'HeadTemp3': 45, 'Cylinder1_3': 46, 'Cylinder2_3': 47,
    'Cylinder3_3': 48, 'ScrewTemp3': 49, 'ScrewSpeed1': 50, 'ScrewSpeed2': 51, 'ScrewSpeed3': 52,
    'FeedRollRatio1': 53, 'FeedRollRatio2': 54, 'FeedRollRatio3': 55, 'Pressure1': 56,
    'Pressure2': 57, 'Pressure3': 58, 'AmMeter': 59, 'AmMeter2': 60, 'AmMeter3': 61,
    'PresetValue': 62, 'ControlValue': 63, 'SpiralPitchSetting': 64, 'SpiralPitchDisplay': 65,
    'SpiralSpeed': 66, 'HoseSpeed': 67, 'UnsmoothSurface': 68, 'MarkingSort': 69,
    'TextMarkingMaterial': 70, 'MarkingColour': 71, 'ChillerWaterTemp': 72, 'CaterpillarGap': 73,
    'TakeUpConveyorSpeed': 74, 'CoolConveyorSpeed': 75, 'CoolConveyorSpeed2': 76, 'ConveyorRatio': 77,
    'ToleranceInner': 78, 'ToleranceOuter': 79, 'TebalInner': 80, 'TebalInnerMiddle': 81,
    'TebalTotal': 82, 'SelisihTebal': 83, 'InnerTarget': 84, 'InnerTol': 85, 'InnerLCL': 86,
    'InnerMin': 87, 'InnerUCL': 88, 'InnerMax': 89, 'ThickTarget': 90, 'ThickTol': 91,
    'ThickLCL': 92, 'ThickMin': 93, 'ThickUCL': 94, 'ThickMax': 95, 'InnerMidTarget': 96,
    'InnerMidTol': 97, 'InnerMidLCL': 98, 'InnerMidMin': 99, 'InnerMidUCL': 100,
    'InnerMidMax': 101, 'TotalTarget': 102, 'TotalTol': 103, 'TotalLCL': 104,
    'TotalMin': 105, 'TotalUCL': 106, 'TotalMax': 107,
}

OUTDATED_DOCS = {
    'SOP/ENG/25/02/12': {'file': r'folder excel\Masterlist SPS CHS 2 Layer DIG.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_CHS2, 'machine': 'CHS 2 Layer'},
    'SOP/PROD': {'file': r'folder excel\Masterlist SPS CHS 2 Layer DIG.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_CHS2, 'machine': 'CHS 2 Layer'},
    'SOP/PROD/HOSE/24/10/054': {'file': r'folder excel\Masterlist SPS CHS 2 Layer DIG.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_CHS2, 'machine': 'CHS 2 Layer'},
    'SOP/PROD/HOSE/24/10/058': {'file': r'folder excel\Masterlist SPS CHS 3 Layer DIG.xlsx', 'sheet': 'Master 1', 'map': MAPPING_CHS3, 'machine': 'CHS 3 Layer'},
    'SOP/PROD/HOSE/24/10/070': {'file': r'folder excel\Masterlist SPS CHS 3 Layer DIG.xlsx', 'sheet': 'Master 1', 'map': MAPPING_CHS3, 'machine': 'CHS 3 Layer'},
    'SOP/PROD/HOSE/24/10/071': {'file': r'folder excel\Masterlist SPS CHS 3 Layer DIG.xlsx', 'sheet': 'Master 1', 'map': MAPPING_CHS3, 'machine': 'CHS 3 Layer'},
    'SOP/PROD/HOSE/EPR/24/10/13': {'file': r'folder excel\Masterlist SPS Double Layer.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_DL, 'machine': 'Double Layer'},
    'SOP/PROD/HOSE/EPR/24/10/16': {'file': r'folder excel\Masterlist SPS Double Layer.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_DL, 'machine': 'Double Layer'},
    'SOP/PROD/HOSE/EPR/24/10/17': {'file': r'folder excel\Masterlist SPS Double Layer.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_DL, 'machine': 'Double Layer'},
    'SOP/PROD/HOSE/EPR/24/10/22': {'file': r'folder excel\Masterlist SPS Double Layer.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_DL, 'machine': 'Double Layer'},
    'SOP/PROD/HOSE/EPR/24/10/24': {'file': r'folder excel\Masterlist SPS Double Layer.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_DL, 'machine': 'Double Layer'},
    'SOP/PROD/HOSE/EPR/24/10/27': {'file': r'folder excel\Masterlist SPS Double Layer.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_DL, 'machine': 'Double Layer'},
    'SOP/PROD/HOSE/EPR/24/10/28': {'file': r'folder excel\Masterlist SPS Double Layer.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_DL, 'machine': 'Double Layer'},
    'SOP/PROD/HOSE/EPR/24/10/37': {'file': r'folder excel\Masterlist SPS Double Layer.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_DL, 'machine': 'Double Layer'},
    'VI-SOP-PROD-129': {'file': r'folder excel\Masterlist SPS CHS 2 Layer DIG.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_CHS2, 'machine': 'CHS 2 Layer'},
    'VI-SOP-PROD-130': {'file': r'folder excel\Masterlist SPS CHS 2 Layer DIG.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_CHS2, 'machine': 'CHS 2 Layer'},
    'VI-SOP-PROD-132': {'file': r'folder excel\Masterlist SPS CHS 2 Layer DIG.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_CHS2, 'machine': 'CHS 2 Layer'},
    'VI-SOP-PROD-133': {'file': r'folder excel\Masterlist SPS CHS 2 Layer DIG.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_CHS2, 'machine': 'CHS 2 Layer'},
    'VI-SOP-PROD-134': {'file': r'folder excel\Masterlist SPS CHS 2 Layer DIG.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_CHS2, 'machine': 'CHS 2 Layer'},
    'VI-SOP-PROD-135': {'file': r'folder excel\Masterlist SPS CHS 2 Layer DIG.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_CHS2, 'machine': 'CHS 2 Layer'},
    'VI-SOP-PROD-137': {'file': r'folder excel\Masterlist SPS CHS 2 Layer DIG.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_CHS2, 'machine': 'CHS 2 Layer'},
    'VI-SOP-PROD-385': {'file': r'folder excel\Masterlist SPS CHS 2 Layer DIG.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_CHS2, 'machine': 'CHS 2 Layer'},
    'VI-SOP-PROD-459': {'file': r'folder excel\Masterlist SPS CHS 2 Layer DIG.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_CHS2, 'machine': 'CHS 2 Layer'},
    'VI-SOP-PROD-461': {'file': r'folder excel\Masterlist SPS CHS 2 Layer DIG.xlsx', 'sheet': 'Parameter Setting', 'map': MAPPING_CHS2, 'machine': 'CHS 2 Layer'},
}

def v(ws, row, col):
    val = ws.cell(row=row, column=col).value
    s = str(val).strip() if val is not None else None
    return s if s and s != '-' else None

print("Mengumpulkan data terbaru dari Excel...")
updates_to_make = {}

# Group by file to optimize openpyxl
files_to_open = set(d['file'] for d in OUTDATED_DOCS.values())

for f in files_to_open:
    print(f"Membaca {f}...")
    wb = openpyxl.load_workbook(f, data_only=True)
    
    docs_in_this_file = {k: v for k, v in OUTDATED_DOCS.items() if v['file'] == f}
    
    for doc_num, info in docs_in_this_file.items():
        ws = wb[info['sheet']]
        mapping = info['map']
        
        # doc_col is 4
        doc_col = 4
        
        highest_rev = -1
        latest_row_data = None
        
        for r in range(5, ws.max_row + 1):
            doc_val = v(ws, r, doc_col)
            if doc_val == doc_num:
                # find rev col dynamically from mapping, or hardcode based on map
                rev_col_idx = mapping.get('RevisionNumber', 7)
                rev_val = v(ws, r, rev_col_idx)
                try:
                    rev_num = int(rev_val) if rev_val else 0
                except:
                    rev_num = 0
                
                if rev_num > highest_rev:
                    highest_rev = rev_num
                    row_data = {'IsActive': '1', 'Machine': info['machine'], 'MachineCode': info['machine']}
                    for field, col_idx in mapping.items():
                        row_data[field] = v(ws, r, col_idx)
                    latest_row_data = row_data
                    
        if latest_row_data:
            updates_to_make[doc_num] = latest_row_data
            
    wb.close()

print(f"Berhasil mengumpulkan data terbaru untuk {len(updates_to_make)} dokumen.")

driver = [d for d in pyodbc.drivers() if 'SQL Server' in d][-1]
conn_str = f"DRIVER={{{driver}}};SERVER={SQL_SERVER};DATABASE={SQL_DB};UID={SQL_USER};PWD={SQL_PASSWORD};TrustServerCertificate=yes;Encrypt=no;"
conn = pyodbc.connect(conn_str)
cur = conn.cursor()

def parse_pm_value(s):
    if not s: return None, None, None
    s = s.replace(',', '.').strip()
    m = re.search(r'([0-9.]+)\s*[±+]\s*([0-9.]+)', s)
    if m:
        base = float(m.group(1))
        tol = float(m.group(2))
        return round(base - tol, 4), round(base, 4), round(base + tol, 4)
    return None, None, None

columns_to_parse = [
    'ToleranceInner', 'ToleranceOuter', 'TebalInner', 'TebalOuter', 'TebalTotal', 
    'TebalInnerMiddle', 'Nipple', 'TubeDie', 'CoverDie', 'MiddleDie', 'SpacerDie', 
    'ADistance', 'MeshDim1', 'MeshDim2', 'MeshDim3', 'HeadTemp1', 'HeadTemp2', 'HeadTemp3',
    'Cylinder1_1', 'Cylinder1_2', 'Cylinder1_3', 'Cylinder2_1', 'Cylinder2_2', 'Cylinder2_3',
    'Cylinder3_1', 'Cylinder3_2', 'Cylinder3_3', 'ScrewTemp1', 'ScrewTemp2', 'ScrewTemp3',
    'ScrewSpeed1', 'ScrewSpeed2', 'ScrewSpeed3', 'Pressure1', 'Pressure2', 'Pressure3',
    'HoseSpeed', 'TakeUpConveyorSpeed', 'ChillerWaterTemp', 'Feed1', 'Feed2', 'Feed3',
    'FeedRollRatio1', 'FeedRollRatio2', 'FeedRollRatio3', 'CaterpillarGap', 'SpiralSpeed',
    'AmMeter', 'AmMeter2', 'AmMeter3', 'PresetValue', 'ControlValue', 'SpiralPitchSetting',
    'SpiralPitchDisplay', 'CoolConveyorSpeed', 'CoolConveyorSpeed2', 'ConveyorRatio',
    'SelisihTebal', 'PitchYarn', 'DancerPosition', 'OdSensor', 'CuttingSpeed', 'ToleranceSpiralPitch'
]

print("Mulai memperbarui database...")
success_count = 0
for doc_num, data in updates_to_make.items():
    try:
        # Delete old
        cur.execute("DELETE FROM SpsNoDocs WHERE DocumentNumber = ?", (doc_num,))
        
        # Insert new
        fields = list(data.keys())
        placeholders = ', '.join(['?' for _ in fields])
        col_names = ', '.join([f'[{f}]' for f in fields])
        values = [data[f] for f in fields]
        
        sql = f'INSERT INTO SpsNoDocs ({col_names}) VALUES ({placeholders})'
        cur.execute(sql, values)
        
        # Parse Min/Asli/Max
        updates = []
        for col in columns_to_parse:
            if col in data and data[col]:
                mn, asli, mx = parse_pm_value(data[col])
                if mn is not None:
                    updates.append(f"[{col}_Min] = {mn}")
                    updates.append(f"[{col}_Asli] = {asli}")
                    updates.append(f"[{col}_Max] = {mx}")
                    
        if updates:
            update_sql = f"UPDATE SpsNoDocs SET {', '.join(updates)} WHERE DocumentNumber = ?"
            cur.execute(update_sql, (doc_num,))
            
        success_count += 1
        print(f"  [+] Berhasil update: {doc_num} (Rev {data.get('RevisionNumber')})")
    except Exception as e:
        print(f"  [!] Gagal update {doc_num}: {e}")

conn.commit()
conn.close()
print(f"\\nSelesai! {success_count} dokumen berhasil di-update.")
