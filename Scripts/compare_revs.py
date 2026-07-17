import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')
EXCEL_FILE = r"folder excel\Masterlist SPS CHS 2 Layer DIG.xlsx"
SHEET_NAME = "Parameter Setting"

MAPPING_CHS2 = {
    'No': 2, 'DocumentNumber': 3, 'RevisionNumber': 6, 'Customer': 7, 'RevisionDate': 8,
    'Formulasi': 9, 'HoseType': 10, 'Dimensi': 11, 'Material': 12, 'InnerTube': 13,
    'OuterCover': 14, 'UseLimitsInner': 15, 'UseLimitsOuter': 16, 'Yarn': 17, 'PitchYarn': 18,
    'TensionYarnInner': 19, 'TensionYarnOuter': 20, 'Nipple': 21, 'TubeDie': 22, 'MiddleDie': 23,
    'CoverDie': 24, 'SpacerDie': 25, 'ADistance': 26, 'MeshDim1': 27, 'MeshScreen1': 28,
    'MeshDim2': 29, 'MeshScreen2': 30, 'HeadTemp1': 31, 'Cylinder1_1': 32, 'Cylinder2_1': 33,
    'Cylinder3_1': 34, 'ScrewTemp1': 35, 'HeadTemp2': 36, 'Cylinder1_2': 37, 'Cylinder2_2': 38,
    'Cylinder3_2': 39, 'ScrewTemp2': 40, 'ScrewSpeed1': 41, 'ScrewSpeed2': 42, 'FeedRollRatio1': 43,
    'FeedRollRatio2': 44, 'Pressure1': 45, 'Pressure2': 46, 'CurrentValue': 47, 'AmMeter': 48,
    'AmMeter2': 49, 'PresetValue': 50, 'ControlValue': 51, 'SpiralPitchSetting': 52,
    'SpiralPitchDisplay': 53, 'SpiralSpeed': 54, 'HoseSpeed': 55, 'UnsmoothSurface': 56,
    'MarkingSort': 57, 'TextMarkingMaterial': 58, 'MarkingColour': 59, 'ChillerWaterTemp': 60,
    'DancerPosition': 61, 'CaterpillarGap': 62, 'TakeUpConveyorSpeed': 63, 'CoolConveyorSpeed': 64,
    'ConveyorRatio': 65, 'ToleranceInner': 66, 'ToleranceOuter': 67, 'TebalInner': 68,
    'TebalTotal': 69, 'SelisihTebal': 70, 'InnerTarget': 71, 'InnerTol': 72, 'InnerLCL': 73,
    'InnerMin': 74, 'InnerUCL': 75, 'InnerMax': 76, 'ThickTarget': 77, 'ThickTol': 78,
    'ThickLCL': 79, 'ThickMin': 80, 'ThickUCL': 81, 'ThickMax': 82, 'TotalTarget': 83,
    'TotalTol': 84, 'TotalLCL': 85, 'TotalMin': 86, 'TotalUCL': 87, 'TotalMax': 88,
}

# The actual mapping used in the previous fix script:
MAPPING_ACTUAL = {
    'No': 3, 'DocumentNumber': 4, 'RevisionNumber': 7, 'Customer': 8, 'RevisionDate': 9,
    'Formulasi': 10, 'HoseType': 11, 'Dimensi': 12, 'Material': 13, 'InnerTube': 14,
    'OuterCover': 15, 'UseLimitsInner': 16, 'UseLimitsOuter': 17, 'Yarn': 18, 'PitchYarn': 19,
    'TensionYarnInner': 20, 'TensionYarnOuter': 21, 'Nipple': 22, 'TubeDie': 23, 'MiddleDie': 24,
    'CoverDie': 25, 'SpacerDie': 26, 'ADistance': 27, 'MeshDim1': 28, 'MeshScreen1': 29,
    'MeshDim2': 30, 'MeshScreen2': 31, 'HeadTemp1': 32, 'Cylinder1_1': 33, 'Cylinder2_1': 34,
    'Cylinder3_1': 35, 'ScrewTemp1': 36, 'HeadTemp2': 37, 'Cylinder1_2': 38, 'Cylinder2_2': 39,
    'Cylinder3_2': 40, 'ScrewTemp2': 41, 'ScrewSpeed1': 42, 'ScrewSpeed2': 43, 'FeedRollRatio1': 44,
    'FeedRollRatio2': 45, 'Pressure1': 46, 'Pressure2': 47, 'CurrentValue': 48, 'AmMeter': 49,
    'AmMeter2': 50, 'PresetValue': 51, 'ControlValue': 52, 'SpiralPitchSetting': 53,
    'SpiralPitchDisplay': 54, 'SpiralSpeed': 55, 'HoseSpeed': 56, 'UnsmoothSurface': 57,
    'MarkingSort': 58, 'TextMarkingMaterial': 59, 'MarkingColour': 60, 'ChillerWaterTemp': 61,
    'DancerPosition': 62, 'CaterpillarGap': 63, 'TakeUpConveyorSpeed': 64, 'CoolConveyorSpeed': 65,
    'ConveyorRatio': 66, 'ToleranceInner': 67, 'ToleranceOuter': 68, 'TebalInner': 69,
    'TebalTotal': 70, 'SelisihTebal': 71, 'InnerTarget': 72, 'InnerTol': 73, 'InnerLCL': 74,
    'InnerMin': 75, 'InnerUCL': 76, 'InnerMax': 77, 'ThickTarget': 78, 'ThickTol': 79,
    'ThickLCL': 80, 'ThickMin': 81, 'ThickUCL': 82, 'ThickMax': 83, 'TotalTarget': 84,
    'TotalTol': 85, 'TotalLCL': 86, 'TotalMin': 87, 'TotalUCL': 88, 'TotalMax': 89,
}

def v(ws, row, col):
    val = ws.cell(row=row, column=col).value
    s = str(val).strip() if val is not None else None
    return s if s and s != '-' else None

wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
ws = wb[SHEET_NAME]
doc_col = 4

revs = {}
for r in range(6, ws.max_row + 1):
    if v(ws, r, doc_col) == 'VI-SOP-PROD-131':
        rev_val = v(ws, r, 7)
        try:
            rev_num = int(rev_val) if rev_val else 0
        except:
            rev_num = 0
            
        row_data = {}
        for field, col_idx in MAPPING_ACTUAL.items():
            row_data[field] = v(ws, r, col_idx)
        revs[rev_num] = row_data

wb.close()

if 2 in revs and 5 in revs:
    rev2 = revs[2]
    rev5 = revs[5]
    print(f"Perbandingan VI-SOP-PROD-131 (Rev 2 vs Rev 5):\\n")
    differences = []
    for field in MAPPING_ACTUAL.keys():
        val2 = rev2.get(field)
        val5 = rev5.get(field)
        if val2 != val5:
            differences.append((field, val2, val5))
            print(f"- Kolom '{field}':")
            print(f"    Rev 2 (lama) : {val2}")
            print(f"    Rev 5 (baru) : {val5}")
    if not differences:
        print("Tidak ada perbedaan lain.")
else:
    print("Tidak bisa menemukan Rev 2 dan Rev 5.")
