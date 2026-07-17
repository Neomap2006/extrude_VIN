import pyodbc
import openpyxl

driver = [d for d in pyodbc.drivers() if 'SQL Server' in d][-1]
conn = pyodbc.connect(f'DRIVER={{{driver}}};SERVER=10.14.149.34,1433;DATABASE=prd_extrude_hose;UID=usrvelasto;PWD=H1s@na2025!!;TrustServerCertificate=yes;Encrypt=no;')
cur = conn.cursor()

# Get all column names of SpsNoDocs
cur.execute("SELECT COLUMN_NAME FROM INFORMATION_SCHEMA.COLUMNS WHERE TABLE_NAME = 'SpsNoDocs'")
columns = [row[0] for row in cur.fetchall()]
print(f"Kolom di SpsNoDocs: {', '.join(columns)}")

# Identify revision column
rev_col = None
date_col = None
for col in columns:
    if 'rev' in col.lower():
        rev_col = col
    if 'date' in col.lower():
        date_col = col

if rev_col:
    # 1. Check database
    cur.execute(f"SELECT {rev_col}, Machine FROM SpsNoDocs WHERE DocumentNumber = 'VI-SOP-PROD-131'")
    row = cur.fetchone()
    if row:
        print(f'Database -> Document: VI-SOP-PROD-131, Revision: {row[0]}, Machine: {row[1]}')
    else:
        print('Database -> Tidak ditemukan')

# 2. Check Excel
print('\\nMengecek di Excel (Masterlist_Digitalisasi_Combined_Fix.xlsx)...')
wb = openpyxl.load_workbook(r'folder excel\\Masterlist_Digitalisasi_Combined_Fix.xlsx', data_only=True)
ws = wb['CHS 2']
found = []
for r in range(6, ws.max_row + 1):
    doc_val = ws.cell(row=r, column=4).value
    if doc_val and str(doc_val).strip() == 'VI-SOP-PROD-131':
        rev = ws.cell(row=r, column=7).value
        date_val = ws.cell(row=r, column=9).value
        found.append({'row': r, 'rev': rev, 'date': date_val})

for f in found:
    print(f"Excel Baris {f['row']} -> Revision: {f['rev']}, Date: {f['date']}")
