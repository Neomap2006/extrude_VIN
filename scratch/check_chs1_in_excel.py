import openpyxl
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

# List of document numbers we need to find
target_docs = [
    'SOP/PROD/HOSE/SPS/1/16',
    'SOP/PROD/HOSE/SPS/1/24',
    'SOP/PROD/HOSE/SPS/1/27',
    'SOP/PROD/HOSE/SPS/1/37',
    'SOP/PROD/HOSE/SPS/1/38',
    'SOP/PROD/HOSE/SPS/24/10/16',
    'SOP/PROD/HOSE/SPS/24/10/17',
    'SOP/PROD/HOSE/SPS/24/10/24',
    'SOP/PROD/HOSE/SPS/24/10/27',
    'SOP/PROD/HOSE/SPS/24/10/37',
    'SOP/PROD/HOSE/SPS/24/12/38',
    'SOP/PROD/HOSE/SPS/3/16',
]

results = {}

# ==============================
# File 1: Masterlist SPS Double Layer_Digitalisasi.xlsx
# ==============================
print("=== Cek File: Masterlist SPS Double Layer_Digitalisasi.xlsx (sheet: Parameter Setting) ===")
try:
    wb = openpyxl.load_workbook(r'd:\Project\extrude\folder excel\Masterlist SPS Double Layer_Digitalisasi.xlsx', data_only=True)
    ws = wb['Parameter Setting']
    
    # Based on earlier analysis: EPR PARAMETER at col 49 (idx 48), MC at col 52 (idx 51)
    # Also there may be an SPS document number column - let's find it by checking row 5
    header_row = 5
    col_map = {}
    for cell in ws[header_row]:
        v = str(cell.value or "").strip().upper()
        if v:
            col_map[v] = cell.column
    
    print("Kolom yang ditemukan di row 5:", {k: v for k, v in col_map.items() if k in ['MC', 'EPR PARAMETER', 'SPS PARAMETER', 'NO. DOCUMENT', 'DOCUMENT']})
    
    for row in ws.iter_rows(min_row=6, values_only=True):
        for cell_val in row:
            val = str(cell_val or "").strip()
            if val in target_docs:
                # Get full row to find MC
                row_list = list(row)
                mc = str(row_list[51] if len(row_list) > 51 else "").strip()
                results[val] = f"DL Excel: {mc}"
                print(f"  FOUND: {val} -> MC={mc}")
except Exception as e:
    print("Error:", e)

# ==============================
# File 2: Masterlist_Digitalisasi_Combined_Fix.xlsx
# ==============================
print("\n=== Cek File: Masterlist_Digitalisasi_Combined_Fix.xlsx (sheet: DL) ===")
try:
    import pandas as pd
    df = pd.read_excel(r'd:\Project\extrude\folder excel\Masterlist_Digitalisasi_Combined_Fix.xlsx', sheet_name='DL', header=None)
    
    header_idx = -1
    for i in range(15):
        row_str = ' '.join([str(x).upper() for x in df.iloc[i].values])
        if 'MC' in row_str:
            header_idx = i
            break
    
    df.columns = [str(c).strip().upper() for c in df.iloc[header_idx]]
    df = df.iloc[header_idx+1:]
    
    doc_cols = [c for c in df.columns if 'DOCUMENT' in c or 'DOC' in c]
    mc_cols = [c for c in df.columns if c == 'MC']
    
    if doc_cols and mc_cols:
        for _, row in df.iterrows():
            doc = str(row[doc_cols[0]]).strip()
            mc = str(row[mc_cols[0]]).strip()
            if doc in target_docs:
                print(f"  FOUND: {doc} -> MC={mc}")
                if doc not in results:
                    results[doc] = f"Combined Excel: {mc}"
except Exception as e:
    print("Error:", e)

# ==============================
# Summary
# ==============================
print("\n=== SUMMARY ===")
for doc in target_docs:
    if doc in results:
        print(f"  {doc} -> {results[doc]}")
    else:
        print(f"  {doc} -> TIDAK DITEMUKAN di Excel")
