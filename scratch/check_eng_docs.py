import pyodbc
import os
import sys
import io
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

d = [d for d in pyodbc.drivers() if 'SQL Server' in d][-1]
conn = pyodbc.connect(f'DRIVER={d};SERVER=10.14.149.34,1433;DATABASE=prd_extrude_hose;UID=usrvelasto;PWD=H1s@na2025!!;TrustServerCertificate=yes')
cr = conn.cursor()

# Cek semua dokumen SOP/ENG di database
cr.execute("SELECT No, DocumentNumber, Machine FROM SpsNoDocs WHERE DocumentNumber LIKE 'SOP/ENG%' ORDER BY DocumentNumber")
eng_rows = cr.fetchall()
print(f"=== SOP/ENG di database: {len(eng_rows)} baris ===")
for r in eng_rows:
    print(f"  No={r[0]} | {r[1]} | Machine={r[2]}")

# Cari di semua file Excel
excel_dir = r'd:\Project\extrude\folder excel'
print(f"\n=== Mencari SOP/ENG di semua file Excel ===")
target_docs = [r[1] for r in eng_rows]

for fname in os.listdir(excel_dir):
    if not fname.endswith('.xlsx'):
        continue
    fpath = os.path.join(excel_dir, fname)
    try:
        wb = openpyxl.load_workbook(fpath, data_only=True, read_only=True)
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            for row in ws.iter_rows(values_only=True):
                for cell in row:
                    val = str(cell or '').strip()
                    if val in target_docs:
                        print(f"  FOUND [{fname}] Sheet={sheet}: {val}")
        wb.close()
    except Exception as e:
        print(f"  Error {fname}: {e}")
